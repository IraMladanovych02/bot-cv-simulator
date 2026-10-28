from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.parsers.base import (
    BaseDocumentParser,
    DocumentFormat,
    EmptyDocumentError,
    InvalidDocumentError,
    ParsedDocument
)
from app.parsers.text import normalize_document_text


class DocxDocumentParser(BaseDocumentParser):
    """Extract text from paragraphs and tables in DOCX documents."""

    supported_extensions = frozenset({".docx"})

    def parse(self, *, content: bytes, filename: str) -> ParsedDocument:
        """Extract and normalize text from a DOCX document."""

        try:
            document = Document(BytesIO(content))
        except (PackageNotFoundError, ValueError, KeyError) as error:
            raise InvalidDocumentError("The DOCX document is corrupted or invalid.") from error
        text_parts: list[str] = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                row_values = [
                    normalize_document_text(cell.text)
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if row_values:
                    text_parts.append(" | ".join(row_values))

        text = normalize_document_text("\n".join(text_parts))
        if not text:
            raise EmptyDocumentError("The DOCX document does not contain readable text.")

        return ParsedDocument(
            filename=filename,
            document_format=DocumentFormat.DOCX,
            text=text,
        )